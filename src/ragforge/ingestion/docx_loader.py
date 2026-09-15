from pathlib import Path

from docx import Document as DOCXDocument

from ragforge.ingestion.document import Document


class DOCXDocumentLoader:
    """
    Loader for Microsoft Word .docx documents.

    Extracts paragraph text and converts the document
    into the canonical RAGForge Document representation.
    """

    def load(
        self,
        path: str | Path,
    ) -> Document:
        """
        Load a DOCX document from disk.
        """

        path = Path(path)

        if not path.exists():
            raise FileNotFoundError(
                f"Document not found: {path}"
            )

        if not path.is_file():
            raise ValueError(
                f"Path is not a file: {path}"
            )

        if path.suffix.lower() != ".docx":
            raise ValueError(
                "DOCXDocumentLoader only supports .docx files."
            )

        docx = DOCXDocument(str(path))

        paragraphs = [
            paragraph.text
            for paragraph in docx.paragraphs
            if paragraph.text.strip()
        ]

        text = "\n\n".join(
            paragraphs
        ).strip()

        if not text:
            raise ValueError(
                f"DOCX contains no extractable text: {path}"
            )

        return Document(
            text=text,
            source=str(path),
            metadata={
                "file_type": "docx",
                "paragraph_count": len(paragraphs),
            },
        )